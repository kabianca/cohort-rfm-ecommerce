"""`make data`: fetch the UCI Online Retail workbook and flatten it to CSV.

The dataset is distributed as one .xlsx inside a zip. Spark reads CSV, not
Excel, so the conversion happens once here and bronze ingests the CSV. The
conversion is deliberately dumb: every cell becomes its text form, dates in
ISO format, an empty string for a missing value. Typing is silver's job, and
doing it here would hide the traps bronze exists to preserve.

The archive's checksum is pinned. If UCI ever re-publishes the file, every
number in the README changes, and this is where that would surface instead
of in a quiet drift of the segments.
"""

import argparse
import csv
import hashlib
import io
import sys
import urllib.request
import zipfile
from datetime import datetime
from pathlib import Path

from ancora.config import RAW_DIR

URL = "https://archive.ics.uci.edu/static/public/352/online+retail.zip"
SHA256 = "f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b"
ARCHIVE = RAW_DIR / "online_retail.zip"
CSV_PATH = RAW_DIR / "online_retail.csv"


def fetch(force: bool = False) -> Path:
    ARCHIVE.parent.mkdir(parents=True, exist_ok=True)
    if ARCHIVE.exists() and not force:
        print(f"{ARCHIVE} already present, skipping download")
    else:
        print(f"downloading {URL}")
        with urllib.request.urlopen(URL, timeout=120) as response:
            ARCHIVE.write_bytes(response.read())
    digest = hashlib.sha256(ARCHIVE.read_bytes()).hexdigest()
    if digest != SHA256:
        sys.exit(
            f"checksum mismatch for {ARCHIVE}\n  expected {SHA256}\n  got      {digest}\n"
            "UCI may have re-published the dataset; the numbers in the README would no longer hold."
        )
    return ARCHIVE


def cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def convert(archive: Path, force: bool = False) -> Path:
    if CSV_PATH.exists() and not force:
        print(f"{CSV_PATH} already present, skipping conversion")
        return CSV_PATH

    # Imported here so the rest of the package does not need openpyxl at all.
    import openpyxl

    with zipfile.ZipFile(archive) as zf:
        (member,) = [n for n in zf.namelist() if n.endswith(".xlsx")]
        workbook = openpyxl.load_workbook(io.BytesIO(zf.read(member)), read_only=True)

    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [cell_to_text(h) for h in next(rows)]

    count = 0
    with CSV_PATH.open("w", newline="") as out:
        writer = csv.writer(out)
        writer.writerow(header)
        for row in rows:
            writer.writerow(cell_to_text(v) for v in row)
            count += 1
    print(f"wrote {count} rows to {CSV_PATH}")
    return CSV_PATH


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--force", action="store_true", help="re-download and re-convert")
    args = parser.parse_args(argv)
    convert(fetch(args.force), args.force)


if __name__ == "__main__":
    main()
